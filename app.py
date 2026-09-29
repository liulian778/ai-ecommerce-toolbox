import streamlit as st
import requests
import pandas as pd
from io import BytesIO
from datetime import datetime
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


st.set_page_config(
    page_title="AI 电商工具箱",
    page_icon="🛒",
    layout="wide"
)


DEFAULT_PROVIDER_CONFIGS = {
    "智谱 BigModel": {
        "api_url": "https://open.bigmodel.cn/api/paas/v4/chat/completions",
        "model": "glm-4-flash",
    },
    "DeepSeek": {
        "api_url": "https://api.deepseek.com/chat/completions",
        "model": "deepseek-chat",
    },
    "豆包 / 火山方舟": {
        "api_url": "https://ark.cn-beijing.volces.com/api/v3/chat/completions",
        "model": "",
    },
    "阿里云百炼 / 通义千问": {
        "api_url": "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions",
        "model": "qwen-plus",
    },
    "Kimi / Moonshot": {
        "api_url": "https://api.moonshot.cn/v1/chat/completions",
        "model": "",
    },
    "硅基流动 SiliconFlow": {
        "api_url": "https://api.siliconflow.cn/v1/chat/completions",
        "model": "",
    },
    "自定义 OpenAI 兼容": {
        "api_url": "",
        "model": "",
    },
}


EXAMPLE_DATA = {
    "product_name": "奶白色通勤保温杯",
    "product_type": "保温杯 / 水杯",
    "target_user": "上班族、学生、通勤人群",
    "target_platform": "小红书、抖音、TikTok Shop",
    "price_range": "59-129元",
    "style": "奶白色、简约高级、通勤风",
    "selling_points": "轻便、防漏、保温时间长、适合办公室和通勤携带",
    "competitor_info": "竞品多强调保温和容量，但视觉风格普通，缺少生活方式场景表达。",
    "user_reviews": "部分用户担心杯子漏水、杯盖不好清洗、容量不够、外观不够高级。",
}


OUTPUT_COLUMN_ORDER = [
    "name",
    "point",
    "user",
    "platform",
    "price",
    "style",
    "title",
    "main_copy",
    "ai_status",
    "product_position",
    "target_user_ai",
    "core_selling_point_ai",
    "main_image_direction",
    "detail_page_structure",
    "short_video_points",
    "risk_tips",
    "main_image_prompt_ai",
    "scene_image_prompt_ai",
    "background_prompt_ai",
    "negative_prompt",
    "main_copy_1",
    "main_copy_2",
    "main_copy_3",
    "full_image_prompt",
    "background_prompt",
    "prompt",
    "raw_ai_answer",
]


def ask_ai(api_url, api_key, model, prompt):
    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {api_key}",
    }

    payload = {
        "model": model,
        "messages": [
            {
                "role": "system",
                "content": "你是一名资深电商商品策划顾问、电商视觉设计师、短视频带货策划师和 AI 生图提示词专家。",
            },
            {
                "role": "user",
                "content": prompt,
            },
        ],
        "temperature": 0.7,
    }

    try:
        response = requests.post(
            api_url,
            headers=headers,
            json=payload,
            timeout=90,
        )

        if response.status_code != 200:
            return False, f"API 请求失败，状态码：{response.status_code}，返回内容：{response.text}"

        data = response.json()
        answer = data["choices"][0]["message"]["content"]
        return True, answer

    except Exception as e:
        return False, str(e)


def build_diagnosis_prompt(
    product_name,
    product_type,
    target_user,
    target_platform,
    price_range,
    style,
    selling_points,
    competitor_info,
    user_reviews,
    market_type,
    output_language,
    output_mode,
):
    detail_requirement = "请输出详细版，内容尽量完整，有分析、有判断、有可执行建议。" if output_mode == "详细版" else "请输出精简版，重点清晰，适合快速查看。"

    language_requirement = {
        "中文": "请全部使用中文输出。",
        "English": "Please output everything in English.",
        "中英双语": "请使用中英双语输出，每个重点先写中文，再写英文。",
    }[output_language]

    return f"""
你是一名资深电商商品策划顾问，请根据下面的信息，生成一份专业的商品诊断报告。

【商品基础信息】
商品名称：{product_name}
商品类型：{product_type}
目标人群：{target_user}
目标平台：{target_platform}
价格带：{price_range}
设计风格：{style}
主要卖点：{selling_points}

【补充信息】
竞品信息：{competitor_info}
用户评论 / 用户顾虑：{user_reviews}

【市场类型】
{market_type}

【输出要求】
{language_requirement}
{detail_requirement}

请按照下面结构输出：

一、商品机会判断
二、目标用户画像
三、核心卖点提炼
四、主图视觉方向
五、短视频带货方向
六、直播 / 详情页卖点
七、用户顾虑与评论区引导
八、跨境电商适配建议
九、合规风险提醒
十、最终可执行建议
"""


def build_markdown_report(product_name, answer):
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    return f"""# AI 电商商品诊断报告：{product_name}

生成时间：{now}

---

{answer}
"""


def build_txt_report(product_name, answer):
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    return f"""AI 电商商品诊断报告：{product_name}

生成时间：{now}

{answer}
"""


def fill_example():
    for key, value in EXAMPLE_DATA.items():
        st.session_state[key] = value


def remember_api_settings(prefix, provider, api_url, model, api_key):
    st.session_state[f"{prefix}_saved_provider"] = provider
    st.session_state[f"{prefix}_saved_api_url"] = api_url
    st.session_state[f"{prefix}_saved_model"] = model
    st.session_state[f"{prefix}_saved_api_key"] = api_key


def get_provider_defaults(prefix, provider):
    default_config = DEFAULT_PROVIDER_CONFIGS[provider]
    saved_provider = st.session_state.get(f"{prefix}_saved_provider")

    if saved_provider == provider:
        api_url = st.session_state.get(f"{prefix}_saved_api_url", default_config["api_url"])
        model = st.session_state.get(f"{prefix}_saved_model", default_config["model"])
    else:
        api_url = default_config["api_url"]
        model = default_config["model"]

    api_key = st.session_state.get(f"{prefix}_saved_api_key", "")

    return api_url, model, api_key


def make_title(row):
    name = str(row.get("name", "")).strip()
    user = str(row.get("user", "")).strip()
    point = str(row.get("point", "")).strip()

    parts = [name, point, f"适合{user}" if user else ""]
    return "｜".join([p for p in parts if p])


def make_main_copy(row):
    point = str(row.get("point", "")).strip()
    user = str(row.get("user", "")).strip()

    if point and user:
        return f"{point}，适合{user}"
    if point:
        return point
    return "突出商品核心卖点"


def make_basic_prompt(row, image_prompt_language="中文"):
    name = str(row.get("name", "")).strip()
    point = str(row.get("point", "")).strip()
    user = str(row.get("user", "")).strip()
    platform = str(row.get("platform", "")).strip()
    price = str(row.get("price", "")).strip()
    style = str(row.get("style", "")).strip()

    if image_prompt_language == "中文":
        image_language_requirement = "主图提示词、场景图提示词、背景图提示词、负面提示词必须使用中文输出。"
    elif image_prompt_language == "English":
        image_language_requirement = "主图提示词、场景图提示词、背景图提示词、负面提示词 must be written in English."
    else:
        image_language_requirement = "主图提示词、场景图提示词、背景图提示词、负面提示词必须使用中英双语输出，先中文后英文。"

    return f"""
你是一名资深电商商品策划顾问、电商视觉设计师、AI 生图提示词专家。
请根据商品信息，生成结构化商品策划内容。

商品名称：{name}
核心卖点：{point}
目标人群：{user}
目标平台：{platform}
价格带：{price}
视觉风格：{style}

重要要求：
1. 必须严格按照指定字段输出。
2. 每一行格式必须是：字段名|||内容
3. 字段名和内容之间必须使用三个竖线：|||
4. 不要输出 JSON。
5. 不要输出 Markdown。
6. 不要添加多余解释。
7. 每个字段只输出一行。
8. 内容尽量简洁，适合放进 Excel 单元格。
9. 生图提示词要具体，适合 Midjourney、即梦、豆包、通义万相、可灵等工具。
10. 不要出现真实品牌名，不要出现侵权元素。
11. {image_language_requirement}

请严格按以下字段输出，每行一个字段，||| 后填写内容：

商品定位|||
目标用户|||
核心卖点|||
主图方向|||
详情页结构|||
短视频卖点|||
风险提醒|||
主图提示词|||
场景图提示词|||
背景图提示词|||
负面提示词|||
主图文案1|||
主图文案2|||
主图文案3|||
"""


def make_full_image_prompt(row):
    name = str(row.get("name", "")).strip()
    point = str(row.get("point", "")).strip()
    user = str(row.get("user", "")).strip()
    platform = str(row.get("platform", "")).strip()
    style = str(row.get("style", "")).strip()

    return (
        f"电商主图摄影，商品为{name}，突出{point}，适合{user}，"
        f"整体风格为{style}，产品主体清晰，画面干净高级，柔和自然光，"
        f"低饱和色调，左侧或上方预留文字空间，适合{platform}主图，"
        f"1:1构图，不要出现文字、logo、水印、乱码。"
    )


def make_background_prompt(row):
    name = str(row.get("name", "")).strip()
    style = str(row.get("style", "")).strip()
    user = str(row.get("user", "")).strip()

    return (
        f"电商主图背景图，不要出现{name}本体，"
        f"画面风格为{style}，适合{user}，干净高级，柔和自然光，"
        f"低饱和色调，适合后期合成商品，不要文字、logo、水印。"
    )


def process_excel(uploaded_file, image_prompt_language="中文"):
    df = pd.read_excel(uploaded_file)

    required_columns = ["name", "point", "user", "platform", "price", "style"]
    missing_columns = [col for col in required_columns if col not in df.columns]

    if missing_columns:
        return None, f"Excel 缺少这些列：{', '.join(missing_columns)}"

    result_df = df.copy()
    result_df["title"] = result_df.apply(make_title, axis=1)
    result_df["main_copy"] = result_df.apply(make_main_copy, axis=1)
    result_df["prompt"] = result_df.apply(
        lambda row: make_basic_prompt(row, image_prompt_language),
        axis=1
    )
    result_df["full_image_prompt"] = result_df.apply(make_full_image_prompt, axis=1)
    result_df["background_prompt"] = result_df.apply(make_background_prompt, axis=1)

    return result_df, None


def order_columns(df):
    ordered = [col for col in OUTPUT_COLUMN_ORDER if col in df.columns]
    remaining = [col for col in df.columns if col not in ordered]
    return df[ordered + remaining]


def dataframe_to_excel_bytes(df):
    output = BytesIO()

    export_df = order_columns(df.copy())

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export_df.to_excel(writer, index=False, sheet_name="result")

        worksheet = writer.sheets["result"]
        worksheet.freeze_panes = "A2"

        header_fill = PatternFill("solid", fgColor="D9EAF7")
        header_font = Font(bold=True, color="000000")
        wrap_alignment = Alignment(wrap_text=True, vertical="top")

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap_alignment

        for idx, column_cells in enumerate(worksheet.columns, start=1):
            column_letter = get_column_letter(idx)
            header = worksheet.cell(row=1, column=idx).value

            if header in ["prompt", "raw_ai_answer"]:
                width = 30
            elif header and "prompt" in str(header):
                width = 45
            elif header in [
                "product_position",
                "target_user_ai",
                "core_selling_point_ai",
                "main_image_direction",
                "detail_page_structure",
                "short_video_points",
                "risk_tips",
            ]:
                width = 35
            else:
                width = 18

            worksheet.column_dimensions[column_letter].width = width

    return output.getvalue()


def create_products_template_bytes():
    template_df = pd.DataFrame([
        {
            "name": "保温杯",
            "point": "保温防漏",
            "user": "上班族",
            "platform": "淘宝/小红书/抖音",
            "price": "59-129元",
            "style": "简约高级通勤风",
        },
        {
            "name": "桌面收纳盒",
            "point": "分类收纳",
            "user": "学生",
            "platform": "淘宝/拼多多/小红书",
            "price": "19-59元",
            "style": "奶油风桌面整理",
        },
        {
            "name": "护腰靠垫",
            "point": "久坐护腰",
            "user": "办公室人群",
            "platform": "淘宝/小红书/抖音",
            "price": "79-199元",
            "style": "舒适办公简约质感",
        },
    ])

    return dataframe_to_excel_bytes(template_df)


def parse_ai_fields(answer):
    fields = {
        "product_position": "",
        "target_user_ai": "",
        "core_selling_point_ai": "",
        "main_image_direction": "",
        "detail_page_structure": "",
        "short_video_points": "",
        "risk_tips": "",
        "main_image_prompt_ai": "",
        "scene_image_prompt_ai": "",
        "background_prompt_ai": "",
        "negative_prompt": "",
        "main_copy_1": "",
        "main_copy_2": "",
        "main_copy_3": "",
    }

    mapping = {
        "商品定位": "product_position",
        "目标用户": "target_user_ai",
        "核心卖点": "core_selling_point_ai",
        "主图方向": "main_image_direction",
        "详情页结构": "detail_page_structure",
        "短视频卖点": "short_video_points",
        "风险提醒": "risk_tips",
        "主图提示词": "main_image_prompt_ai",
        "场景图提示词": "scene_image_prompt_ai",
        "背景图提示词": "background_prompt_ai",
        "负面提示词": "negative_prompt",
        "主图文案1": "main_copy_1",
        "主图文案2": "main_copy_2",
        "主图文案3": "main_copy_3",
    }

    for line in answer.splitlines():
        line = line.strip()
        if "|||" not in line:
            continue

        key, value = line.split("|||", 1)
        key = key.strip()
        value = value.strip()

        if key in mapping:
            fields[mapping[key]] = value

    return fields


def batch_generate_ai_plans(result_df, api_url, api_key, model, max_items):
    ai_statuses = []

    product_position_list = []
    target_user_ai_list = []
    core_selling_point_ai_list = []
    main_image_direction_list = []
    detail_page_structure_list = []
    short_video_points_list = []
    risk_tips_list = []
    main_image_prompt_ai_list = []
    scene_image_prompt_ai_list = []
    background_prompt_ai_list = []
    negative_prompt_list = []
    main_copy_1_list = []
    main_copy_2_list = []
    main_copy_3_list = []
    raw_answer_list = []

    limited_df = result_df.head(max_items).copy()
    total = len(limited_df)

    progress_bar = st.progress(0)
    status_text = st.empty()

    for index, row in limited_df.iterrows():
        product_name = str(row.get("name", "")).strip()
        prompt = str(row.get("prompt", "")).strip()

        status_text.write(f"正在生成第 {len(ai_statuses) + 1} / {total} 个商品：{product_name}")

        success, answer = ask_ai(
            api_url=api_url,
            api_key=api_key,
            model=model,
            prompt=prompt
        )

        raw_answer_list.append(answer)

        if success:
            parsed = parse_ai_fields(answer)

            ai_statuses.append("成功")
            product_position_list.append(parsed["product_position"])
            target_user_ai_list.append(parsed["target_user_ai"])
            core_selling_point_ai_list.append(parsed["core_selling_point_ai"])
            main_image_direction_list.append(parsed["main_image_direction"])
            detail_page_structure_list.append(parsed["detail_page_structure"])
            short_video_points_list.append(parsed["short_video_points"])
            risk_tips_list.append(parsed["risk_tips"])
            main_image_prompt_ai_list.append(parsed["main_image_prompt_ai"])
            scene_image_prompt_ai_list.append(parsed["scene_image_prompt_ai"])
            background_prompt_ai_list.append(parsed["background_prompt_ai"])
            negative_prompt_list.append(parsed["negative_prompt"])
            main_copy_1_list.append(parsed["main_copy_1"])
            main_copy_2_list.append(parsed["main_copy_2"])
            main_copy_3_list.append(parsed["main_copy_3"])
        else:
            ai_statuses.append("失败")
            product_position_list.append(answer)
            target_user_ai_list.append("")
            core_selling_point_ai_list.append("")
            main_image_direction_list.append("")
            detail_page_structure_list.append("")
            short_video_points_list.append("")
            risk_tips_list.append("")
            main_image_prompt_ai_list.append("")
            scene_image_prompt_ai_list.append("")
            background_prompt_ai_list.append("")
            negative_prompt_list.append("")
            main_copy_1_list.append("")
            main_copy_2_list.append("")
            main_copy_3_list.append("")

        progress_bar.progress(len(ai_statuses) / total)

    output_df = limited_df.copy()
    output_df["ai_status"] = ai_statuses
    output_df["product_position"] = product_position_list
    output_df["target_user_ai"] = target_user_ai_list
    output_df["core_selling_point_ai"] = core_selling_point_ai_list
    output_df["main_image_direction"] = main_image_direction_list
    output_df["detail_page_structure"] = detail_page_structure_list
    output_df["short_video_points"] = short_video_points_list
    output_df["risk_tips"] = risk_tips_list
    output_df["main_image_prompt_ai"] = main_image_prompt_ai_list
    output_df["scene_image_prompt_ai"] = scene_image_prompt_ai_list
    output_df["background_prompt_ai"] = background_prompt_ai_list
    output_df["negative_prompt"] = negative_prompt_list
    output_df["main_copy_1"] = main_copy_1_list
    output_df["main_copy_2"] = main_copy_2_list
    output_df["main_copy_3"] = main_copy_3_list
    output_df["raw_ai_answer"] = raw_answer_list

    status_text.write("批量生成完成。")

    return order_columns(output_df)


def show_home():
    st.title("AI 电商工具箱")
    st.write("一个面向电商设计师、运营和跨境卖家的 AI 商品策划工具箱。")

    st.info("这个工具箱可以帮助你完成单品诊断、商品策划、主图方向、短视频卖点和 AI 生图提示词生成。")

    col1, col2 = st.columns(2)

    with col1:
        st.subheader("商品诊断工具")
        st.write("适合单个商品深度分析。")
        st.markdown("""
- 判断商品机会
- 分析目标用户
- 提炼核心卖点
- 给出主图方向
- 给出短视频方向
- 生成诊断报告
""")

    with col2:
        st.subheader("商品策划生成工具")
        st.write("适合多个商品批量生成。")
        st.markdown("""
- 上传 Excel 商品表
- 批量生成商品标题
- 批量生成主图文案
- 批量生成 AI 生图提示词
- 下载结构化 Excel
""")

    st.subheader("使用建议")
    st.markdown("""
1. 如果你只有一个商品，先用「商品诊断工具」。
2. 如果你有多个商品，使用「商品策划生成工具」。
3. API Key 只在本次会话中使用，不会写入代码。
4. 批量生成时请不要切换工具、刷新页面或关闭网页。
""")


def show_diagnosis_tool():
    st.title("AI 电商商品诊断工具")
    st.write("适合单个商品深度分析，生成一份商品诊断报告。")

    if st.button("一键填入示例"):
        fill_example()
        st.rerun()

    st.divider()

    st.subheader("1. 商品基础信息")

    col1, col2 = st.columns(2)

    with col1:
        product_name = st.text_input("商品名称", key="product_name")
        product_type = st.text_input("商品类型", key="product_type")
        target_user = st.text_input("目标人群", key="target_user")
        target_platform = st.text_input("目标平台", key="target_platform")

    with col2:
        price_range = st.text_input("价格带", key="price_range")
        style = st.text_input("设计风格", key="style")
        selling_points = st.text_area("主要卖点", height=120, key="selling_points")

    st.subheader("2. 可选补充信息")

    competitor_info = st.text_area(
        "竞品信息",
        height=120,
        key="competitor_info",
        placeholder="可以填写竞品价格、卖点、主图风格、差异点等。没有也可以不填。"
    )

    user_reviews = st.text_area(
        "用户评论 / 用户顾虑",
        height=120,
        key="user_reviews",
        placeholder="可以粘贴用户评论，也可以填写你猜测的用户顾虑。没有也可以不填。"
    )

    st.subheader("3. 市场与输出设置")

    col3, col4, col5 = st.columns(3)

    with col3:
        market_type = st.selectbox(
            "市场类型",
            ["中国电商", "跨境电商", "Amazon", "Etsy", "Shopify", "TikTok Shop"]
        )

    with col4:
        output_language = st.selectbox(
            "输出语言",
            ["中文", "English", "中英双语"]
        )

    with col5:
        output_mode = st.selectbox(
            "输出模式",
            ["详细版", "精简版"]
        )

    st.subheader("4. 选择模型平台")

    provider = st.selectbox(
        "请选择模型平台",
        list(DEFAULT_PROVIDER_CONFIGS.keys()),
        key="diagnosis_provider"
    )

    default_api_url, default_model, saved_api_key = get_provider_defaults("diagnosis", provider)

    api_url = st.text_input(
        "API 地址",
        value=default_api_url,
        key=f"diagnosis_api_url_{provider}"
    )

    model = st.text_input(
        "模型名称",
        value=default_model,
        key=f"diagnosis_model_{provider}"
    )

    api_key = st.text_input(
        "请输入所选平台的 API Key",
        value=saved_api_key,
        type="password",
        key=f"diagnosis_api_key_{provider}"
    )

    if st.button("记住本次诊断工具 API 设置"):
        remember_api_settings("diagnosis", provider, api_url, model, api_key)
        st.success("已记住本次会话的诊断工具 API 设置。")

    st.info("隐私提示：API Key 只在本次网页运行中使用，不会写入代码文件。不要截图或公开你的真实 Key。")

    if st.button("生成商品诊断报告"):
        if not product_name:
            st.warning("请至少填写商品名称。")
            return

        if not api_url:
            st.warning("请填写 API 地址。")
            return

        if not model:
            st.warning("请填写模型名称。")
            return

        if not api_key:
            st.warning("请填写 API Key。")
            return

        prompt = build_diagnosis_prompt(
            product_name,
            product_type,
            target_user,
            target_platform,
            price_range,
            style,
            selling_points,
            competitor_info,
            user_reviews,
            market_type,
            output_language,
            output_mode,
        )

        remember_api_settings("diagnosis", provider, api_url, model, api_key)

        with st.spinner("AI 正在生成商品诊断报告，请稍等..."):
            success, result = ask_ai(api_url, api_key, model, prompt)

        if success:
            st.session_state["diagnosis_answer"] = result
            st.session_state["diagnosis_product_name"] = product_name
            st.success("生成成功！")
        else:
            st.error(result)

    if "diagnosis_answer" in st.session_state:
        answer = st.session_state["diagnosis_answer"]
        report_product_name = st.session_state.get("diagnosis_product_name", "商品")

        st.subheader("5. 商品诊断报告")
        st.markdown(answer)

        markdown_report = build_markdown_report(report_product_name, answer)
        txt_report = build_txt_report(report_product_name, answer)

        st.info("下载格式说明：Markdown（.md）适合保留标题、列表等结构；TXT（.txt）适合直接打开、复制和转发。")

        col_down1, col_down2 = st.columns(2)

        with col_down1:
            st.download_button(
                label="下载 Markdown 诊断报告",
                data=markdown_report,
                file_name=f"{report_product_name}_商品诊断报告.md",
                mime="text/markdown"
            )

        with col_down2:
            st.download_button(
                label="下载 TXT 诊断报告",
                data=txt_report,
                file_name=f"{report_product_name}_商品诊断报告.txt",
                mime="text/plain"
            )


def show_planner_tool():
    st.title("AI 电商商品策划生成工具")
    st.write("适合多个商品批量生成标题、主图文案、策划方向和 AI 生图提示词。")

    st.subheader("1. 上传商品表")

    st.info("Excel 表格需要包含这些列：name、point、user、platform、price、style")

    image_prompt_language = st.selectbox(
        "生图提示词语言",
        ["中文", "English", "中英双语"]
    )

    template_excel = create_products_template_bytes()

    st.download_button(
        label="下载商品表模板 products_template.xlsx",
        data=template_excel,
        file_name="products_template.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    uploaded_file = st.file_uploader(
        "请上传 products.xlsx",
        type=["xlsx"]
    )

    if uploaded_file is not None:
        result_df, error = process_excel(uploaded_file, image_prompt_language)

        if error:
            st.error(error)
            return

        st.session_state["planner_result_df"] = result_df
        st.session_state["planner_image_prompt_language"] = image_prompt_language

    if "planner_result_df" in st.session_state:
        result_df = st.session_state["planner_result_df"]

        st.success(f"基础内容生成成功！共识别 {len(result_df)} 个商品。")

        st.subheader("2. 基础生成结果预览")
        st.dataframe(result_df, use_container_width=True)

        excel_data = dataframe_to_excel_bytes(result_df)

        st.download_button(
            label="下载基础版 output.xlsx",
            data=excel_data,
            file_name="output.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        st.subheader("3. 批量生成 AI 策划与生图提示词")

        max_items = st.number_input(
            "本次最多处理多少个商品",
            min_value=1,
            max_value=len(result_df),
            value=min(10, len(result_df)),
            step=1
        )

        provider = st.selectbox(
            "请选择模型平台",
            list(DEFAULT_PROVIDER_CONFIGS.keys()),
            key="planner_provider"
        )

        default_api_url, default_model, saved_api_key = get_provider_defaults("planner", provider)

        api_url = st.text_input(
            "API 地址",
            value=default_api_url,
            key=f"planner_api_url_{provider}"
        )

        model = st.text_input(
            "模型名称",
            value=default_model,
            key=f"planner_model_{provider}"
        )

        api_key = st.text_input(
            "请输入所选平台的 API Key",
            value=saved_api_key,
            type="password",
            key=f"planner_api_key_{provider}"
        )

        if st.button("记住本次策划工具 API 设置"):
            remember_api_settings("planner", provider, api_url, model, api_key)
            st.success("已记住本次会话的策划工具 API 设置。")

        st.warning("重要提示：批量生成过程中请不要切换左侧工具、刷新页面或关闭网页，否则当前任务会中断。")

        if st.button("批量生成 AI 策划"):
            if not api_url:
                st.warning("请填写 API 地址。")
                return

            if not model:
                st.warning("请填写模型名称。")
                return

            if not api_key:
                st.warning("请填写 API Key。")
                return

            remember_api_settings("planner", provider, api_url, model, api_key)

            output_ai_df = batch_generate_ai_plans(
                result_df=result_df,
                api_url=api_url,
                api_key=api_key,
                model=model,
                max_items=int(max_items)
            )

            st.session_state["planner_output_ai_df"] = output_ai_df
            st.success("批量 AI 策划生成完成！")

    if "planner_output_ai_df" in st.session_state:
        output_ai_df = st.session_state["planner_output_ai_df"]

        st.subheader("4. AI 策划结果预览")
        st.dataframe(output_ai_df, use_container_width=True)

        output_ai_excel = dataframe_to_excel_bytes(output_ai_df)

        st.download_button(
            label="下载完整版 output_ai.xlsx",
            data=output_ai_excel,
            file_name="output_ai.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    st.subheader("5. 表格字段说明")

    st.markdown("""
| 字段名 | 含义 | 示例 |
|---|---|---|
| name | 商品名称 | 保温杯 |
| point | 核心卖点 | 保温防漏 |
| user | 目标用户 | 上班族 |
| platform | 目标平台 | 淘宝/小红书/抖音 |
| price | 价格带 | 59-129元 |
| style | 视觉风格 | 简约高级通勤风 |
""")


st.sidebar.title("AI 电商工具箱")

tool = st.sidebar.radio(
    "请选择工具",
    [
        "首页 / 使用说明",
        "商品诊断工具",
        "商品策划生成工具",
    ]
)

st.sidebar.divider()
st.sidebar.caption("当前版本：产品化原型版")
st.sidebar.caption("适合：电商设计师 / 运营 / 跨境卖家 / AI 工具学习者")
st.sidebar.warning("批量生成时请不要切换工具、刷新页面或关闭网页，否则任务会中断。")

if tool == "首页 / 使用说明":
    show_home()

elif tool == "商品诊断工具":
    show_diagnosis_tool()

elif tool == "商品策划生成工具":
    show_planner_tool()